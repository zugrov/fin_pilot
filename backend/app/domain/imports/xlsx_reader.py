import zipfile
from io import BytesIO

from openpyxl import load_workbook

from app.domain.imports.contracts import CASH_COLUMNS, ImportFormatError, ImportLine
from app.domain.imports.import_meta import validate_import_header
from app.domain.imports.normalize import cell_text, normalize_row


def read_xlsx(content: bytes, *, max_rows: int, max_uncompressed: int) -> list[ImportLine]:
    _guard_size(content, max_uncompressed)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ImportFormatError("unsupported_format", "Файл XLSX не читается") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise ImportFormatError("unsupported_format", "В книге нет листа")
        iterator = sheet.iter_rows()
        try:
            header_row = next(iterator)
        except StopIteration as exc:
            raise ImportFormatError("unsupported_format", "Пустой лист") from exc
        header = [cell_text(cell.value)[0] for cell in header_row]
        while header and header[-1] == "":
            header.pop()
        columns = validate_import_header(tuple(header), CASH_COLUMNS)
        lines: list[ImportLine] = []
        for index, cells in enumerate(iterator, start=2):
            if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                continue
            if len(lines) >= max_rows:
                raise ImportFormatError("unsupported_format", "Слишком много строк")
            raw, extra = _cells(cells, columns)
            line = normalize_row(raw, index, date_style="YYYY-MM-DD")
            if extra:
                line = ImportLine(
                    line.row_number,
                    line.external_id,
                    line.wallet_code,
                    line.business_date,
                    line.direction,
                    line.amount,
                    line.currency,
                    line.counterparty_inn,
                    line.counterparty_name,
                    line.description,
                    line.document_number,
                    line.category_code,
                    line.direction_code,
                    tuple(dict.fromkeys([*line.errors, *extra])),
                    line.raw,
                )
            lines.append(line)
        return lines
    finally:
        workbook.close()


def _cells(cells, columns: tuple[str, ...]) -> tuple[dict[str, str], tuple[str, ...]]:
    raw: dict[str, str] = {}
    errors: list[str] = []
    for pos, column in enumerate(columns):
        if column not in CASH_COLUMNS and pos >= len(CASH_COLUMNS):
            cell = cells[pos] if pos < len(cells) else None
            text = "" if cell is None else cell_text(cell.value)[0]
            raw[column] = text
            continue
        cell = cells[pos] if pos < len(cells) else None
        if cell is None:
            raw[column] = ""
            continue
        if getattr(cell, "data_type", None) == "f" and column == "amount":
            raw[column] = ""
            errors.append("formula_without_value")
            continue
        text, error = cell_text(cell.value)
        if error == "amount_as_float" and column != "amount":
            text = ""
            error = "bad_cell"
        raw[column] = text
        if error and column == "amount":
            errors.append(error)
        elif error:
            errors.append(error)
    return raw, tuple(dict.fromkeys(errors))


def _guard_size(content: bytes, limit: int) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ImportFormatError("unsupported_format", "Файл не является XLSX") from exc
    if total > limit:
        raise ImportFormatError("unsupported_format", "Распакованный XLSX больше лимита")
